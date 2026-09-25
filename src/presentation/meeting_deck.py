"""Build outputs/presentation/Workhuman_Taxonomy_Meeting_Progress.pptx from verified metrics only."""
from __future__ import annotations

import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from src.utils.common import OUTPUTS

M = json.loads((OUTPUTS / "tables" / "sprint1_metrics.json").read_text(encoding="utf-8"))
TC = M["table_counts"]
OUT = OUTPUTS / "presentation" / "Workhuman_Taxonomy_Meeting_Progress.pptx"
FIG = OUTPUTS / "figures"

INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x8A, 0x89, 0x84)
BLUE, ORANGE, AQUA = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A)
LIGHT, WHITE, GREEN, AMBER, GREY = (RGBColor(0xF0, 0xEF, 0xEC), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x0C, 0xA3, 0x0C),
                                    RGBColor(0xC9, 0x85, 0x00), RGBColor(0x9A, 0x99, 0x94))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def pct(x):
    return f"{100 * x:.1f}%"


def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"
    return tb


def box(slide, x, y, w, h, s, fill=LIGHT, color=INK, size=13, bold=False, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid()
    b.fill.fore_color.rgb = fill
    if line:
        b.line.color.rgb = line
        b.line.width = Pt(1.5)
    else:
        b.line.fill.background()
    tf = b.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold if i == 0 else False, color, "Calibri"
    return b


def arrow(slide, x1, y1, x2, y2, color=GREY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.append(tail)
    return c


def header(slide, title, sub=None, n=None):
    text(slide, 0.6, 0.35, 12, 0.7, title, 28, True)
    if sub:
        text(slide, 0.6, 1.0, 12, 0.5, sub, 15, False, INK2)
    if n:
        text(slide, 12.2, 7.0, 1, 0.4, str(n), 10, False, MUTED, PP_ALIGN.RIGHT)
    text(slide, 0.6, 7.0, 9, 0.4, "IE7945 Workforce Analytics capstone x Workhuman - Taxonomy of Work", 10, False, MUTED)


def badge(slide, x, y, label):
    col = {"COMPLETED": GREEN, "IN PROGRESS": AMBER, "NEXT": GREY}[label]
    box(slide, x, y, 1.45, 0.34, label, fill=col, color=WHITE, size=10, bold=True)


# 1 END GOAL ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Taxonomy of Work", "End goal: connect real labour-market evidence to business decisions", 1)
chain = ["Public Job Postings", "Tasks + Skills", "Task / Skill Taxonomies", "Task-Skill Mapping",
         "Business Initiative -> Tasks -> Skills"]
for i, c in enumerate(chain):
    y = 1.7 + i * 1.0
    box(s, 3.9, y, 5.5, 0.7, c, fill=BLUE if i in (0, 4) else LIGHT, color=WHITE if i in (0, 4) else INK, size=17, bold=True)
    if i < len(chain) - 1:
        arrow(s, 6.65, y + 0.7, 6.65, y + 1.0)
text(s, 9.8, 1.8, 3.2, 2, "Evidence comes from real postings;\nstandards (O*NET, ESCO) are the reference frame.", 13, False, INK2)

# 2 COMPLETED --------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "A reproducible shared workforce-data foundation", "Sprint 1 - completed and verified", 2)
for i, (v, l) in enumerate([(M["raw_postings"], "raw postings"), (M["canonical_postings"], "canonical postings"),
                            (M["unique_postings"], "unique postings")]):
    text(s, 0.8 + i * 4.1, 1.7, 3.8, 1.0, f"{v:,}", 54, True, BLUE, PP_ALIGN.CENTER)
    text(s, 0.8 + i * 4.1, 2.75, 3.8, 0.5, l, 16, False, INK2, PP_ALIGN.CENTER)
items = [f"O*NET 31.0\n{TC['onet_occupations']:,} occupations | {TC['onet_tasks']:,} tasks",
         f"ESCO v1.2.1\n{TC['esco_occupations']:,} occupations | {TC['esco_skills']:,} skills",
         "DuckDB\none shared database", "Shared data contract\none posting_id for both teams",
         "EDA notebook\nexecuted from the DB", "Automated pipeline\n21/21 acceptance checks"]
for i, it in enumerate(items):
    box(s, 0.8 + (i % 3) * 4.1, 3.7 + (i // 3) * 1.35, 3.8, 1.1, it, size=14, bold=True)
text(s, 0.8, 6.55, 12, 0.4, "Unique count corrected 4,037 -> 4,029 after an HTML-decoding fix revealed 8 hidden duplicates "
     "(docs/sprint1_bugfix_log.md). Raw and canonical counts unchanged.", 10, False, MUTED)

# 3 HOW DATA FITS ----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "How the data fits together", "Job postings = empirical evidence | O*NET & ESCO = reference standards", 3)
box(s, 0.7, 1.9, 2.6, 0.8, f"NYC Open Data Jobs\n{M['raw_by_source']['nyc_jobs']:,} raw", fill=BLUE, color=WHITE, size=14, bold=True)
box(s, 0.7, 3.3, 2.6, 0.8, f"Arbeitnow API\n{M['raw_by_source']['arbeitnow']:,} raw", fill=ORANGE, color=WHITE, size=14, bold=True)
box(s, 4.3, 2.5, 3.0, 1.0, f"Canonical Job Corpus\n{M['canonical_postings']:,} postings", size=15, bold=True, line=BLUE)
arrow(s, 3.3, 2.3, 4.3, 2.9)
arrow(s, 3.3, 3.7, 4.3, 3.1)
box(s, 4.8, 4.0, 2.0, 0.55, "posting_id", fill=INK, color=WHITE, size=14, bold=True)
arrow(s, 5.8, 3.5, 5.8, 4.0)
box(s, 8.3, 1.9, 2.0, 0.8, "TASKS", fill=LIGHT, size=16, bold=True)
box(s, 8.3, 4.6, 2.0, 0.8, "SKILLS", fill=LIGHT, size=16, bold=True)
arrow(s, 6.8, 4.27, 8.3, 2.3)
arrow(s, 6.8, 4.27, 8.3, 5.0)
box(s, 10.9, 1.9, 2.0, 0.8, f"O*NET 31.0\n{TC['onet_tasks']:,} tasks", fill=AQUA, color=WHITE, size=13, bold=True)
box(s, 10.9, 4.6, 2.0, 0.8, f"ESCO v1.2.1\n{TC['esco_skills']:,} skills", fill=AQUA, color=WHITE, size=13, bold=True)
arrow(s, 10.3, 2.3, 10.9, 2.3)
arrow(s, 10.3, 5.0, 10.9, 5.0)
text(s, 0.7, 5.9, 12, 0.9, "Every downstream artefact (task, skill, mapping) joins back to the same posting_id and to a preserved,\n"
     "checksummed raw record - so both teams work on identical evidence.", 14, False, INK2)

# 4 EDA -------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "What EDA taught us", "Corpus quality and representation must be measured before taxonomy construction", 4)
nyc = next(r for r in M["by_source"] if r["source"] == "nyc_jobs")
facts = [(pct(M["duplicate_rate"]), "overall duplicates"), (pct(nyc["duplicate_rate"]), "NYC duplicates\n(internal + external re-posts)"),
         (pct(M["language_mix"]["en"]["pct"]), "English"), (pct(M["language_mix"]["de"]["pct"]), "German"),
         ("0%", "industry available\nfrom current sources")]
for i, (v, l) in enumerate(facts):
    text(s, 0.6, 1.65 + i * 1.0, 2.2, 0.6, v, 30, True, BLUE if i < 2 else (ORANGE if i < 4 else INK2))
    text(s, 2.8, 1.75 + i * 1.0, 2.6, 0.8, l, 13, False, INK2)
s.shapes.add_picture(str(FIG / "duplicates_by_source.png"), Inches(5.6), Inches(1.6), width=Inches(7.2))
s.shapes.add_picture(str(FIG / "language_mix.png"), Inches(6.6), Inches(4.25), height=Inches(2.6))

# 5 DATA CONTRACT ----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Independent development now; integration without rework later", "Shared data contract v1.0 (docs/data_contract.md)", 5)
for i, f in enumerate(["posting_id", "source + source_job_id", "raw record (preserved)", "normalized fields\n(NULL = not available)"]):
    box(s, 0.7, 1.7 + i * 1.05, 3.4, 0.85, f, fill=INK if i == 0 else LIGHT, color=WHITE if i == 0 else INK, size=14, bold=True)
box(s, 5.2, 1.9, 3.3, 1.1, "TASK TEAM\nposting_id -> atomic tasks", fill=BLUE, color=WHITE, size=15, bold=True)
box(s, 5.2, 3.9, 3.3, 1.1, "SKILL TEAM\nposting_id -> atomic skills", fill=ORANGE, color=WHITE, size=15, bold=True)
arrow(s, 4.1, 2.1, 5.2, 2.45)
arrow(s, 4.1, 2.1, 5.2, 4.45)
box(s, 9.6, 2.9, 3.2, 1.1, "LATER\nposting_id\ntasks <-> skills", fill=AQUA, color=WHITE, size=15, bold=True)
arrow(s, 8.5, 2.45, 9.6, 3.3)
arrow(s, 8.5, 4.45, 9.6, 3.6)
text(s, 0.7, 6.1, 12, 0.6, "Rules: UTF-8, ISO-8601 dates, ISO language codes, source-provided vs *_derived fields, duplicates flagged not deleted.", 13, False, INK2)

# 6 ADVANCED PROTOTYPE -----------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Advanced prototype direction", "Model disagreement as a quality signal - not an LLM answer as ground truth", 6)
steps = [("REAL POSTING (deterministic sample, section detection)", "COMPLETED"),
         ("Gemini + Groq independent extraction (cached, schema-validated)", "IN PROGRESS"),
         ("Evidence-preserving gates (verbatim span must exist in posting)", "COMPLETED"),
         ("Task / skill statements + cross-model agreement", "IN PROGRESS"),
         ("O*NET / ESCO candidate alignment", "NEXT"),
         ("Human review where models disagree", "NEXT")]
for i, (st, lab) in enumerate(steps):
    y = 1.6 + i * 0.85
    box(s, 0.8, y, 8.2, 0.62, st, size=14, bold=True)
    badge(s, 9.2, y + 0.14, lab)
    if i < len(steps) - 1:
        arrow(s, 4.9, y + 0.62, 4.9, y + 0.85)
text(s, 10.85, 1.6, 2.3, 4.5, "COMPLETED = built and tested end to end on a 4-posting smoke test.\n\n"
     "IN PROGRESS = 200-posting experiment paused mid-run; responses cached, results not yet validated.\n\n"
     "No experiment metrics are shown until the run completes.", 12, False, INK2)

# 7 COLLABORATION ----------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Collaboration path", "One shared foundation feeds both teams; the joint work reuses their outputs", 7)
box(s, 4.4, 1.55, 4.5, 1.25, "SHARED FOUNDATION\ncanonical corpus | posting IDs\ndata contract | O*NET + ESCO", fill=INK, color=WHITE, size=14, bold=True)
box(s, 0.7, 3.4, 5.0, 1.25, "TASK TEAM\nresponsibilities -> tasks -> clusters\n-> O*NET evaluation", fill=BLUE, color=WHITE, size=14, bold=True)
box(s, 7.6, 3.4, 5.0, 1.25, "SKILL TEAM\nrequirements -> skills -> clusters\n-> ESCO evaluation", fill=ORANGE, color=WHITE, size=14, bold=True)
box(s, 4.4, 5.3, 4.5, 1.25, "JOINT\ntask-skill mapping\ninitiative -> tasks -> skills", fill=AQUA, color=WHITE, size=14, bold=True)
arrow(s, 5.6, 2.8, 3.2, 3.4)
arrow(s, 7.7, 2.8, 10.1, 3.4)
arrow(s, 3.2, 4.65, 5.6, 5.3)
arrow(s, 10.1, 4.65, 7.7, 5.3)

# 8 STATUS ----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
header(s, "Foundation complete; moving from data acquisition to evidence extraction", "Current status", 8)
cols = [("DONE", GREEN, ["Corpus (NYC + Arbeitnow)", "O*NET 31.0 loaded", "ESCO v1.2.1 loaded", "DuckDB database",
                          "Shared data contract", "EDA notebook", "Reproducible pipeline (21/21)"]),
        ("IN PROGRESS", AMBER, ["Gemini/Groq provider layer - built, live-tested", "Deterministic 200-posting sample - built",
                                 "Section detection + evidence gates - built, tested", "Dual-model extraction run - paused (cached)",
                                 "Empty gold-100 set + annotation guidelines - drafted"]),
        ("NEXT", GREY, ["Finish evidence-based extraction prototype", "Gemini/Groq disagreement experiment",
                        "Human-labelled 100-posting set", "O*NET/ESCO candidate matching",
                        "Embeddings/clustering (next sprint)", "Joint task-skill integration later"])]
for i, (t, col, its) in enumerate(cols):
    x = 0.6 + i * 4.2
    box(s, x, 1.55, 3.9, 0.6, t, fill=col, color=WHITE, size=16, bold=True)
    text(s, x + 0.1, 2.35, 3.8, 4.4, "\n".join("- " + it for it in its), 14, False, INK)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
